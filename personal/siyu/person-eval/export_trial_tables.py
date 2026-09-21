from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
ROOT = Path(__file__).resolve().parent
SRC = ROOT / "person-eval-trials.json"
OUT = ROOT / "person-eval-trials-summary.xlsx"
OUT_FALLBACK = ROOT / "person-eval-summary.xlsx"
LABELS = ("fight", "fall", "smoking", "weapon")
LABEL_ZH = {"fight": "打架", "fall": "摔倒", "smoking": "抽烟", "weapon": "持械"}
CLAUDE_MARKERS = ("claude", "fable", "opus", "sonnet", "haiku")

sys.path.insert(0, str(ROOT))
from eval_ten_clips import CLIPS  # noqa: E402


def pct(value):
    return None if value is None else round(float(value), 4)


def num(value, digits=3):
    return None if value is None else round(float(value), digits)


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin
            if cell.row > 1:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, min_w=10, max_w=48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_w, len(value) + 2))
        ws.column_dimensions[letter].width = width


def predicted(row: dict, label: str) -> str:
    if row.get("error"):
        return "ERR"
    for item in row.get("judgments") or ():
        if item.get("event_label") == label:
            return "是" if item.get("detected") else "否"
    return "缺"


def yes_no(flag) -> str:
    return "是" if flag else "否"


def is_claude(model: str) -> bool:
    text = (model or "").lower()
    return any(marker in text for marker in CLAUDE_MARKERS)


def clip_file(clip: dict) -> str:
    return clip.get("file") or clip.get("file_glob") or ""


def sample_annotation(expected: dict) -> str:
    positives = [LABEL_ZH[name] for name in LABELS if expected.get(name)]
    return "、".join(positives) if positives else "全否（无打架/摔倒/抽烟/持械）"


def clip_annotation_summary(clip: dict) -> str:
    bits = []
    for sample in clip["samples"]:
        bits.append(f"{sample['timestamp']}s {sample_annotation(sample['expected'])}")
    return "；".join(bits)


def write_cover(workbook: Workbook) -> None:
    cover = workbook.active
    cover.title = "首页"
    cover["A1"] = "安全事件评测说明"
    cover["A1"].font = Font(size=18, bold=True, color="1F4E79")
    cover.merge_cells("A1:G1")
    cover["A2"] = (
        "本表对比同一条 aiclub 渠道上的 gpt-5.5、gemini-3.8-flash、DeepSeek-V4.1-Flash。"
        "Claude 已从对照中移除。每段视频抽 1～2 帧，对打架、摔倒、抽烟、持械四类独立计分。"
    )
    cover["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    cover.merge_cells("A2:G2")
    cover.row_dimensions[2].height = 36

    counts = {name: 0 for name in LABELS}
    negative_frames = 0
    total_frames = 0
    for clip in CLIPS:
        for sample in clip["samples"]:
            total_frames += 1
            if any(sample["expected"].get(name) for name in LABELS):
                for name in LABELS:
                    counts[name] += int(bool(sample["expected"].get(name)))
            else:
                negative_frames += 1
    cover["A3"] = (
        f"共 {len(CLIPS)} 段视频、{total_frames} 个抽帧。"
        f"正样本：打架 {counts['fight']}、摔倒 {counts['fall']}、"
        f"抽烟 {counts['smoking']}、持械 {counts['weapon']}；"
        f"四类全否 {negative_frames} 帧。"
    )
    cover.merge_cells("A3:G3")

    cover["A5"] = "测试视频与标注"
    cover["A5"].font = Font(size=13, bold=True, color="1F4E79")
    header_row = 6
    headers = [
        "片段",
        "文件",
        "来源",
        "画面说明",
        "抽帧标注",
        "打架帧",
        "摔倒帧",
        "抽烟帧",
        "持械帧",
    ]
    for col, title in enumerate(headers, start=1):
        cover.cell(header_row, col, title)
    for index, clip in enumerate(CLIPS, start=header_row + 1):
        expected_by_label = {name: [] for name in LABELS}
        for sample in clip["samples"]:
            for name in LABELS:
                if sample["expected"].get(name):
                    expected_by_label[name].append(f"{sample['timestamp']}s")
        cover.cell(index, 1, clip["id"])
        cover.cell(index, 2, clip_file(clip))
        cover.cell(index, 3, clip.get("source", ""))
        cover.cell(index, 4, clip.get("scene", ""))
        cover.cell(index, 5, clip_annotation_summary(clip))
        cover.cell(index, 6, "、".join(expected_by_label["fight"]) or "无")
        cover.cell(index, 7, "、".join(expected_by_label["fall"]) or "无")
        cover.cell(index, 8, "、".join(expected_by_label["smoking"]) or "无")
        cover.cell(index, 9, "、".join(expected_by_label["weapon"]) or "无")

    last = header_row + len(CLIPS)
    cover.auto_filter.ref = f"A{header_row}:I{last}"
    cover.freeze_panes = f"A{header_row + 1}"
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for col in range(1, 10):
        cell = cover.cell(header_row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in cover.iter_rows(min_row=header_row, max_row=last, min_col=1, max_col=9):
        for cell in row:
            cell.border = thin
            if cell.row > header_row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    cover.row_dimensions[header_row].height = 22
    cover.column_dimensions["A"].width = 26
    cover.column_dimensions["B"].width = 36
    cover.column_dimensions["C"].width = 42
    cover.column_dimensions["D"].width = 40
    cover.column_dimensions["E"].width = 48
    for letter in "FGHI":
        cover.column_dimensions[letter].width = 14


def main() -> None:
    payload = json.loads(SRC.read_text(encoding="utf-8"))
    trials = [trial for trial in payload["trials"] if not is_claude(trial.get("model", ""))]
    workbook = Workbook()
    write_cover(workbook)

    overview = workbook.create_sheet("模型总览")
    overview.append(
        [
            "模型",
            "渠道",
            "有效请求",
            "总请求",
            "标签正确率",
            "正确标签",
            "计分标签",
            "打架召回",
            "摔倒召回",
            "抽烟召回",
            "持械召回",
            "打架精确",
            "摔倒精确",
            "抽烟精确",
            "持械精确",
            "p50秒",
            "p95秒",
            "平均秒",
            "墙钟秒",
            "漏评标签数",
        ]
    )
    for trial in trials:
        labels = trial["labels"]
        overview.append(
            [
                trial["model"],
                trial["provider"],
                trial["valid"],
                trial["requests"],
                pct(trial["label_accuracy"]),
                trial["correct_labels"],
                trial["scored_labels"],
                pct(labels["fight"]["recall"]),
                pct(labels["fall"]["recall"]),
                pct(labels["smoking"]["recall"]),
                pct(labels["weapon"]["recall"]),
                pct(labels["fight"]["precision"]),
                pct(labels["fall"]["precision"]),
                pct(labels["smoking"]["precision"]),
                pct(labels["weapon"]["precision"]),
                num(trial["p50_latency"]),
                num(trial["p95_latency"]),
                num(trial["mean_latency"]),
                num(trial["wall_seconds"], 1),
                sum(labels[name]["missing"] for name in LABELS),
            ]
        )
    percent_cols = {5, 8, 9, 10, 11, 12, 13, 14, 15}
    for row in range(2, overview.max_row + 1):
        for col in percent_cols:
            overview.cell(row, col).number_format = "0.0%"
        for col in (16, 17, 18):
            overview.cell(row, col).number_format = "0.00"
    style_header(overview)
    autosize(overview)

    events = workbook.create_sheet("事件混淆")
    events.append(["模型", "事件", "TP", "FP", "TN", "FN", "漏评", "精确率", "召回率"])
    for trial in trials:
        for name in LABELS:
            matrix = trial["labels"][name]
            events.append(
                [
                    trial["model"],
                    LABEL_ZH[name],
                    matrix["tp"],
                    matrix["fp"],
                    matrix["tn"],
                    matrix["fn"],
                    matrix["missing"],
                    pct(matrix["precision"]),
                    pct(matrix["recall"]),
                ]
            )
    for row in range(2, events.max_row + 1):
        events.cell(row, 8).number_format = "0.0%"
        events.cell(row, 9).number_format = "0.0%"
    style_header(events)
    autosize(events)

    clips = workbook.create_sheet("分片段")
    clips.append(
        [
            "片段",
            "场景",
            "模型",
            "有效帧",
            "错误帧",
            "标签正确率",
            "正确/计分",
            "平均耗时秒",
            "最长耗时秒",
        ]
    )
    for trial in trials:
        for clip in trial["clips"]:
            clips.append(
                [
                    clip["clip_id"],
                    clip.get("scene", ""),
                    trial["model"],
                    clip["valid"],
                    clip["errors"],
                    pct(clip["label_accuracy"]),
                    f"{clip['correct_labels']}/{clip['scored_labels']}",
                    num(clip["mean_latency"]),
                    num(clip["max_latency"]),
                ]
            )
    for row in range(2, clips.max_row + 1):
        clips.cell(row, 6).number_format = "0.0%"
        clips.cell(row, 8).number_format = "0.00"
        clips.cell(row, 9).number_format = "0.00"
    style_header(clips)
    autosize(clips, max_w=56)

    keys: list[tuple] = []
    seen: set[tuple] = set()
    lookups: list[tuple[str, dict]] = []
    for trial in trials:
        index = {}
        for frame in trial["frames"]:
            key = (frame["clip_id"], frame["timestamp"])
            index[key] = frame
            if key not in seen:
                seen.add(key)
                keys.append(key)
        lookups.append((trial["model"], index))

    frames = workbook.create_sheet("逐帧对照")
    header = ["片段", "秒", "标注-打架", "标注-摔倒", "标注-抽烟", "标注-持械"]
    for model, _ in lookups:
        header.extend(
            [
                f"{model}-打架",
                f"{model}-摔倒",
                f"{model}-抽烟",
                f"{model}-持械",
                f"{model}-对错",
                f"{model}-耗时秒",
                f"{model}-画面说明",
                f"{model}-错误",
            ]
        )
    frames.append(header)
    for key in keys:
        sample = next(index[key] for _, index in lookups if key in index)
        expected = sample.get("expected") or {}
        row = [
            key[0],
            key[1],
            yes_no(expected.get("fight")),
            yes_no(expected.get("fall")),
            yes_no(expected.get("smoking")),
            yes_no(expected.get("weapon")),
        ]
        for _, index in lookups:
            frame = index.get(key) or {}
            score = frame.get("score") or {}
            row.extend(
                [
                    predicted(frame, "fight"),
                    predicted(frame, "fall"),
                    predicted(frame, "smoking"),
                    predicted(frame, "weapon"),
                    f"{score.get('correct', 0)}/{score.get('scored', 0)}",
                    num(frame.get("latency_seconds")),
                    (frame.get("scene_summary") or "").replace("\n", " "),
                    frame.get("error") or "",
                ]
            )
        frames.append(row)
    style_header(frames)
    frames.freeze_panes = "C2"
    autosize(frames, min_w=8, max_w=40)

    dest = OUT_FALLBACK
    workbook.save(dest)
    print(dest)


if __name__ == "__main__":
    main()
